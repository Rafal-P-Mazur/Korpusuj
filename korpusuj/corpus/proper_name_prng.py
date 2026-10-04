# -*- coding: utf-8 -*-
"""Lokalny eksperyment PRNG dla nierozstrzygnietych nazw wlasnych.

Buduje indeks SQLite z lokalnych plikow PRNG (XLSX, CSV/TSV lub ZIP
zawierajacy XLSX/CSV), a nastepnie porownuje z nim obserwacje z JSON-u
experiment_proper_name_lemma_audit_v2.py.

Nie laczy sie z Internetem i nie modyfikuje Parquetu ani anotacji korpusu.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sqlite3
import sys
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator

VERSION = "0.2.0-experimental"
SCHEMA_VERSION = 1
EXTERNAL_DECISION = "EXTERNAL_REGISTRY_CANDIDATE"

# Nazwy kolumn roznia sie miedzy wydaniami PRNG. Klasyfikacja jest celowo
# oparta na znormalizowanych nazwach naglowkow, nie na jednej wersji schematu.
COLUMN_HINTS = {
    "canonical": (
        "polska nazwa", "nazwa polska", "nazwa glowna", "nazwa główna",
        "nazwa urzedowa", "nazwa urzędowa", "nazwa zestandaryzowana",
        "egzonim", "nazwa obiektu", "nazwa",
    ),
    "genitive": (
        "dopelniacz", "dopełniacz", "forma dopelniacza", "forma dopełniacza",
        "odmiana dopelniacz", "odmiana dopełniacz", "genitive",
    ),
    "locative": (
        "miejscownik", "forma miejscownika", "odmiana miejscownik", "locative",
    ),
    "adjective": ("przymiotnik", "forma przymiotnika", "adjective"),
    "variant": (
        "wariant", "nazwa wariantowa", "nazwa oboczna", "nazwa historyczna",
        "endonim", "pseudoegzonim", "inne nazwy", "alias",
    ),
    "object_type": (
        "rodzaj obiektu", "typ obiektu", "klasa obiektu", "kategoria obiektu",
        "object type", "feature type",
    ),
    "country": ("panstwo", "państwo", "kraj", "country"),
    "identifier": (
        "identyfikator", "id prng", "identyfikator prng", "prng id", "uuid",
        "lokalny id", "local id",
    ),
    "status": ("status nazwy", "status", "rodzaj nazwy"),
}

CASE_TO_KIND = {"nom": "canonical", "gen": "genitive", "loc": "locative"}
MATCH_STRENGTH = {
    "PRNG_CASE_FORM_CONFIRMED": 5,
    "PRNG_CANONICAL_FORM_CONFIRMED": 4,
    "PRNG_VARIANT_FORM_CONFIRMED": 3,
    "PRNG_OTHER_CASE_FORM": 2,
    "PRNG_NO_MATCH": 0,
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_display(value: Any) -> str:
    text = clean(value).replace("\u00ad", "")
    text = unicodedata.normalize("NFC", text)
    return re.sub(r"\s+", " ", text).strip()


def norm(value: Any) -> str:
    return unicodedata.normalize("NFKC", normalize_display(value)).casefold()


def norm_header(value: Any) -> str:
    text = norm(value)
    text = re.sub(r"[^0-9a-ząćęłńóśźż]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def safe_md(value: Any) -> str:
    return clean(value).replace("|", "\\|").replace("\n", " ")


def json_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def split_values(value: Any) -> list[str]:
    text = normalize_display(value)
    if not text:
        return []
    # Nie dzielimy po przecinku, bo moze nalezec do nazwy. PRNG zwykle uzywa
    # srednikow, pionowych kresek lub nowych wierszy do list wariantow.
    parts = re.split(r"\s*(?:;|\||\r?\n)\s*", text)
    output = []
    seen = set()
    for part in parts:
        part = normalize_display(part)
        if part and norm(part) not in seen:
            seen.add(norm(part))
            output.append(part)
    return output


def score_header(header: str, hint: str) -> int:
    h, q = norm_header(header), norm_header(hint)
    if not h or not q:
        return 0
    if h == q:
        return 100
    if q in h:
        return 70 + min(len(q), 20)
    if h in q:
        return 45 + min(len(h), 20)
    hset, qset = set(h.split()), set(q.split())
    overlap = len(hset & qset)
    return overlap * 15


def classify_columns(headers: list[str]) -> tuple[dict[str, list[int]], dict[str, Any]]:
    mapping: dict[str, list[int]] = defaultdict(list)
    scored = {}
    for idx, header in enumerate(headers):
        best_kind, best_score = "", 0
        for kind, hints in COLUMN_HINTS.items():
            score = max(score_header(header, hint) for hint in hints)
            if score > best_score:
                best_kind, best_score = kind, score
        if best_kind and best_score >= 30:
            mapping[best_kind].append(idx)
            scored[header] = {"kind": best_kind, "score": best_score}
    # Kanoniczna nazwa musi byc ostrozniejsza niz pozostale pola. Jesli wiele
    # kolumn nazwowych zostalo uznanych za canonical, preferujemy najwyzszy wynik.
    canonical = mapping.get("canonical", [])
    if len(canonical) > 1:
        best = max(canonical, key=lambda i: scored.get(headers[i], {}).get("score", 0))
        for idx in canonical:
            if idx != best:
                mapping["variant"].append(idx)
        mapping["canonical"] = [best]
    return dict(mapping), {"headers": headers, "classified": scored}


def iter_csv_rows(handle: io.TextIOBase, source_name: str) -> Iterator[tuple[str, list[str], Iterator[list[Any]]]]:
    sample = handle.read(65536)
    handle.seek(0)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    reader = csv.reader(handle, dialect)
    try:
        headers = [normalize_display(value) for value in next(reader)]
    except StopIteration:
        return
    yield source_name, headers, reader


def iter_xlsx_rows(stream: Any, source_name: str) -> Iterator[tuple[str, list[str], Iterator[list[Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Brak openpyxl. Zainstaluj: pip install openpyxl") from exc
    workbook = load_workbook(stream, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = None
        buffered = []
        # PRNG moze miec kilka wierszy opisowych przed naglowkiem. Szukamy
        # pierwszego wiersza, ktory wyglada jak schemat nazw geograficznych.
        for _ in range(30):
            try:
                row = list(next(rows))
            except StopIteration:
                break
            values = [normalize_display(value) for value in row]
            mapping, _ = classify_columns(values)
            if mapping.get("canonical") and len([x for x in values if x]) >= 2:
                header = values
                break
            buffered.append(row)
        if header is not None:
            yield f"{source_name}::{sheet.title}", header, rows
    workbook.close()


def iter_sources(paths: list[Path]) -> Iterator[tuple[str, list[str], Iterator[list[Any]]]]:
    for path in paths:
        suffix = path.suffix.casefold()
        if suffix == ".xlsx":
            yield from iter_xlsx_rows(path, path.name)
        elif suffix in {".csv", ".tsv", ".txt"}:
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
                yield from iter_csv_rows(handle, path.name)
        elif suffix == ".zip":
            with zipfile.ZipFile(path) as archive:
                for member in archive.infolist():
                    if member.is_dir():
                        continue
                    member_suffix = Path(member.filename).suffix.casefold()
                    if member_suffix == ".xlsx":
                        data = io.BytesIO(archive.read(member))
                        yield from iter_xlsx_rows(data, f"{path.name}::{member.filename}")
                    elif member_suffix in {".csv", ".tsv", ".txt"}:
                        data = archive.read(member)
                        handle = io.StringIO(data.decode("utf-8-sig", errors="replace"))
                        yield from iter_csv_rows(handle, f"{path.name}::{member.filename}")
        else:
            raise RuntimeError(f"Nieobslugiwany format: {path}")


def connect(path: Path) -> sqlite3.Connection:
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    con.execute("PRAGMA journal_mode=WAL")
    con.execute("PRAGMA synchronous=NORMAL")
    return con


def create_schema(con: sqlite3.Connection) -> None:
    con.executescript("""
    DROP TABLE IF EXISTS names;
    DROP TABLE IF EXISTS objects;
    DROP TABLE IF EXISTS sources;
    DROP TABLE IF EXISTS metadata;
    CREATE TABLE metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL);
    CREATE TABLE sources(
        source_id INTEGER PRIMARY KEY,
        source_name TEXT NOT NULL,
        headers_json TEXT NOT NULL,
        mapping_json TEXT NOT NULL,
        row_count INTEGER NOT NULL DEFAULT 0,
        accepted_count INTEGER NOT NULL DEFAULT 0
    );
    CREATE TABLE objects(
        object_id INTEGER PRIMARY KEY,
        source_id INTEGER NOT NULL REFERENCES sources(source_id),
        source_row INTEGER NOT NULL,
        external_id TEXT NOT NULL,
        canonical_name TEXT NOT NULL,
        canonical_norm TEXT NOT NULL,
        object_type TEXT NOT NULL,
        country TEXT NOT NULL,
        status TEXT NOT NULL,
        raw_json TEXT NOT NULL
    );
    CREATE TABLE names(
        name_id INTEGER PRIMARY KEY,
        object_id INTEGER NOT NULL REFERENCES objects(object_id) ON DELETE CASCADE,
        name TEXT NOT NULL,
        name_norm TEXT NOT NULL,
        kind TEXT NOT NULL,
        source_column TEXT NOT NULL
    );
    CREATE INDEX idx_names_norm ON names(name_norm);
    CREATE INDEX idx_names_norm_kind ON names(name_norm,kind);
    CREATE INDEX idx_objects_canonical ON objects(canonical_norm);
    """)


def first_value(row: list[Any], indexes: list[int]) -> str:
    for idx in indexes:
        if idx < len(row):
            value = normalize_display(row[idx])
            if value:
                return value
    return ""


def values_with_columns(row: list[Any], indexes: list[int], headers: list[str]) -> list[tuple[str, str]]:
    output = []
    seen = set()
    for idx in indexes:
        if idx >= len(row):
            continue
        for value in split_values(row[idx]):
            key = norm(value)
            if key and key not in seen:
                seen.add(key)
                output.append((value, headers[idx]))
    return output


def build_index(source_paths: list[Path], index: Path, batch_size: int) -> dict[str, Any]:
    index.parent.mkdir(parents=True, exist_ok=True)
    if index.exists():
        index.unlink()
    con = connect(index)
    create_schema(con)
    stats = Counter()
    diagnostics = []
    next_object_id = 1
    try:
        for source_name, headers, rows in iter_sources(source_paths):
            mapping, diag = classify_columns(headers)
            cursor = con.execute(
                "INSERT INTO sources(source_name,headers_json,mapping_json) VALUES(?,?,?)",
                (source_name, json_text(headers), json_text(mapping)),
            )
            source_id = cursor.lastrowid
            source_stats = Counter()
            if not mapping.get("canonical"):
                diagnostics.append({"source": source_name, "error": "NO_CANONICAL_COLUMN", **diag})
                continue
            object_rows, name_rows = [], []
            for row_no, raw_row in enumerate(rows, 1):
                row = list(raw_row)
                source_stats["rows"] += 1
                canonical = first_value(row, mapping.get("canonical", []))
                if not canonical:
                    continue
                object_id = next_object_id
                next_object_id += 1
                raw = {headers[i]: clean(row[i]) for i in range(min(len(headers), len(row))) if clean(row[i])}
                external_id = first_value(row, mapping.get("identifier", []))
                object_type = first_value(row, mapping.get("object_type", []))
                country = first_value(row, mapping.get("country", []))
                status = first_value(row, mapping.get("status", []))
                object_rows.append((
                    object_id, source_id, row_no, external_id, canonical, norm(canonical),
                    object_type, country, status, json_text(raw),
                ))
                pairs = [("canonical", canonical, headers[mapping["canonical"][0]])]
                for kind in ("genitive", "locative", "adjective", "variant"):
                    pairs.extend((kind, value, column) for value, column in values_with_columns(row, mapping.get(kind, []), headers))
                seen = set()
                for kind, value, column in pairs:
                    key = (norm(value), kind)
                    if not key[0] or key in seen:
                        continue
                    seen.add(key)
                    name_rows.append((object_id, value, key[0], kind, column))
                source_stats["accepted"] += 1
                source_stats["names"] += len(seen)
                if len(object_rows) >= batch_size:
                    con.executemany("INSERT INTO objects VALUES(?,?,?,?,?,?,?,?,?,?)", object_rows)
                    con.executemany(
                        "INSERT INTO names(object_id,name,name_norm,kind,source_column) VALUES(?,?,?,?,?)",
                        name_rows,
                    )
                    con.commit()
                    object_rows.clear()
                    name_rows.clear()
            if object_rows:
                con.executemany("INSERT INTO objects VALUES(?,?,?,?,?,?,?,?,?,?)", object_rows)
                con.executemany(
                    "INSERT INTO names(object_id,name,name_norm,kind,source_column) VALUES(?,?,?,?,?)",
                    name_rows,
                )
            con.execute(
                "UPDATE sources SET row_count=?,accepted_count=? WHERE source_id=?",
                (source_stats["rows"], source_stats["accepted"], source_id),
            )
            con.commit()
            stats.update(source_stats)
            diagnostics.append({"source": source_name, "mapping": mapping, "stats": dict(source_stats), **diag})
            print(
                f"[prng-index] {source_name}: {source_stats['accepted']:,} obiektow, "
                f"{source_stats['names']:,} nazw",
                file=sys.stderr, flush=True,
            )
        metadata = {
            "schema_version": SCHEMA_VERSION,
            "version": VERSION,
            "built_at": datetime.now(timezone.utc).isoformat(),
            "sources": [str(path.resolve()) for path in source_paths],
            "source_sizes": {str(path.resolve()): path.stat().st_size for path in source_paths},
            "diagnostics": diagnostics,
            "stats": dict(stats),
        }
        con.executemany(
            "INSERT INTO metadata(key,value) VALUES(?,?)",
            [(key, json_text(value)) for key, value in metadata.items()],
        )
        con.commit()
        con.execute("PRAGMA optimize")
    finally:
        con.close()
    return {
        "success": True, "experimental": True, "version": VERSION,
        "index": str(index.resolve()), "sources": len(diagnostics), **dict(stats),
    }


def parse_case(tag: str) -> str:
    parts = [part.casefold() for part in clean(tag).split(":")]
    return parts[2] if len(parts) > 2 else ""


def query_name(con: sqlite3.Connection, value: str) -> list[dict[str, Any]]:
    rows = con.execute("""
        SELECT n.name,n.kind,n.source_column,o.object_id,o.external_id,
               o.canonical_name,o.object_type,o.country,o.status,s.source_name
        FROM names n
        JOIN objects o ON o.object_id=n.object_id
        JOIN sources s ON s.source_id=o.source_id
        WHERE n.name_norm=?
        ORDER BY CASE n.kind WHEN 'canonical' THEN 0 WHEN 'genitive' THEN 1
                 WHEN 'locative' THEN 2 WHEN 'variant' THEN 3 ELSE 4 END,
                 o.canonical_norm,o.object_id
    """, (norm(value),)).fetchall()
    return [dict(row) for row in rows]


def load_candidates(path: Path, include_review: bool, min_count: int) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    output = []
    for row in payload.get("rows", []):
        decision = clean(row.get("decision"))
        if decision != EXTERNAL_DECISION and not (include_review and decision == "REVIEW"):
            continue
        if int(row.get("token_count") or 0) < min_count:
            continue
        output.append(row)
    return output


def classify_candidate(row: dict[str, Any], hits: list[dict[str, Any]]) -> dict[str, Any]:
    observed_case = parse_case(clean(row.get("morph")))
    expected_kind = CASE_TO_KIND.get(observed_case)
    model_lemma = clean(row.get("model_lemma"))
    canonical_names = sorted({hit["canonical_name"] for hit in hits}, key=str.casefold)
    expected_hits = [hit for hit in hits if expected_kind and hit["kind"] == expected_kind]
    canonical_hits = [hit for hit in hits if hit["kind"] == "canonical"]
    variant_hits = [hit for hit in hits if hit["kind"] == "variant"]

    # Najpierw wybieramy rodzaj dowodu. Dopiero potem rozstrzygamy, czy dowod
    # oznacza korekte, czy tylko potwierdzenie lematu modelu.
    if expected_hits:
        evidence_kind = "CASE_FORM"
        selected = expected_hits
        confidence = "HIGH"
    elif canonical_hits and observed_case in {"", "nom"}:
        evidence_kind = "CANONICAL_FORM"
        selected = canonical_hits
        confidence = "HIGH"
    elif variant_hits:
        evidence_kind = "ALTERNATIVE_NAME"
        selected = variant_hits
        confidence = "MEDIUM"
    elif hits:
        evidence_kind = "NAME_ONLY"
        selected = hits
        confidence = "LOW"
    else:
        evidence_kind = "NO_MATCH"
        selected = []
        confidence = "NONE"

    selected_canonicals = sorted({hit["canonical_name"] for hit in selected}, key=str.casefold)
    unique = len(selected_canonicals) == 1
    canonical_name = selected_canonicals[0] if unique else ""
    proposed_lemma = ""
    changes_lemma = False
    decision = "NO_ACTION"

    if selected and not unique:
        category = f"PRNG_{evidence_kind}_AMBIGUOUS"
        confidence = "LOW"
        decision = "REVIEW"
    elif evidence_kind in {"CASE_FORM", "CANONICAL_FORM"} and canonical_name:
        if norm(model_lemma) == norm(canonical_name):
            category = "PRNG_MODEL_LEMMA_CONFIRMED"
            decision = "KEEP"
        else:
            category = "PRNG_SAFE_LEMMA_REPAIR"
            proposed_lemma = canonical_name
            changes_lemma = True
            decision = "AUTO_CANDIDATE"
    elif evidence_kind == "ALTERNATIVE_NAME" and canonical_name:
        # Nazwa oboczna identyfikuje obiekt, ale nie uprawnia do zastapienia
        # lematu tokenu nazwa glowna, np. Kobane -> Ajn al-Arab.
        category = "PRNG_ALTERNATIVE_NAME_CONFIRMED"
        decision = "KEEP"
    elif evidence_kind == "NAME_ONLY":
        category = "PRNG_NAME_ONLY_CONFIRMED"
        decision = "REVIEW"
    else:
        category = "PRNG_NO_MATCH"

    return {
        "orth_raw": clean(row.get("orth")),
        "orth_normalized": normalize_display(row.get("orth")),
        "model_lemma": model_lemma,
        "model_lemma_normalized": normalize_display(model_lemma),
        "local_target": clean(row.get("target")),
        "proposed_lemma": proposed_lemma,
        "canonical_name": canonical_name,
        "changes_lemma": changes_lemma,
        "decision": decision,
        "evidence_kind": evidence_kind,
        "morph": clean(row.get("morph")),
        "observed_case": observed_case,
        "ner_raw": clean(row.get("ner_raw")),
        "span_text": clean(row.get("span_text")),
        "source_category": clean(row.get("category")),
        "token_count": int(row.get("token_count") or 0),
        "document_count": int(row.get("document_count") or 0),
        "category": category,
        "confidence": confidence,
        "unique_canonical": unique,
        "all_canonical_names": canonical_names,
        "hits": hits,
        "examples": row.get("examples", [])[:3],
    }


def aggregate(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups = {}
    for result in results:
        target = result["proposed_lemma"] or result["local_target"]
        key = (norm(result["model_lemma"]), norm(target))
        group = groups.setdefault(key, {
            "model_lemma": result["model_lemma"], "target": target,
            "forms": Counter(), "categories": set(), "confidences": set(),
            "token_count": 0, "document_upper_bound": 0,
        })
        group["forms"][result["orth_raw"]] += result["token_count"]
        group["categories"].add(result["category"])
        group["confidences"].add(result["confidence"])
        group["token_count"] += result["token_count"]
        group["document_upper_bound"] += result["document_count"]
    output = []
    for group in groups.values():
        output.append({
            "model_lemma": group["model_lemma"], "target": group["target"],
            "forms": dict(group["forms"].most_common()),
            "categories": sorted(group["categories"]),
            "confidences": sorted(group["confidences"]),
            "token_count": group["token_count"],
            "document_upper_bound": group["document_upper_bound"],
        })
    return sorted(output, key=lambda item: (-item["token_count"], norm(item["model_lemma"])))


def write_report(path: Path, audit: Path, index: Path, results: list[dict[str, Any]], rules: list[dict[str, Any]]) -> None:
    categories = Counter(result["category"] for result in results)
    confidences = Counter(result["confidence"] for result in results)
    repairs = [result for result in results if result.get("changes_lemma")]
    confirmations = [
        result for result in results
        if result.get("category") == "PRNG_MODEL_LEMMA_CONFIRMED"
    ]
    alternative_names = [
        result for result in results
        if result.get("category") == "PRNG_ALTERNATIVE_NAME_CONFIRMED"
    ]
    repair_tokens = sum(result["token_count"] for result in repairs)
    confirmation_tokens = sum(result["token_count"] for result in confirmations)
    lines = [
        "# Lokalny eksperyment PRNG dla nazw własnych", "",
        f"**Wersja:** `{VERSION}`  ",
        f"**Wygenerowano:** `{datetime.now(timezone.utc).isoformat()}`  ",
        f"**Audyt wejściowy:** `{audit}`  ",
        f"**Indeks PRNG:** `{index}`", "",
        "> Eksperyment diagnostyczny. Nie modyfikuje korpusu.", "",
        "## Podsumowanie", "",
        f"- Obserwacje: **{len(results):,}**",
        f"- Rzeczywiste bezpieczne korekty lematu: **{len(repairs):,}**",
        f"- Tokeny ze zmianą lematu: **{repair_tokens:,}**",
        f"- Potwierdzenia już poprawnego lematu modelu: **{len(confirmations):,}**",
        f"- Tokeny z potwierdzonym lematem bez zmiany: **{confirmation_tokens:,}**",
        f"- Potwierdzone nazwy oboczne bez zamiany lematu: **{len(alternative_names):,}**",
        f"- Zagregowane pary lemat modelu → cel: **{len(rules):,}**", "",
        "## Kategorie", "",
    ]
    for key, value in categories.most_common():
        lines.append(f"- `{key}`: **{value:,}**")
    lines += ["", "## Pewność", ""]
    for key, value in confidences.most_common():
        lines.append(f"- `{key}`: **{value:,}**")
    lines += ["", "## Reguły zagregowane", "", "| Lemat modelu | Cel | Formy | Tokeny | Kategorie | Pewność |", "|---|---|---|---:|---|---|"]
    for rule in rules:
        forms = ", ".join(f"{form} ({count})" for form, count in rule["forms"].items())
        lines.append("| " + " | ".join([
            safe_md(rule["model_lemma"]), safe_md(rule["target"]), safe_md(forms),
            f"{rule['token_count']:,}", safe_md(", ".join(rule["categories"])),
            safe_md(", ".join(rule["confidences"])),
        ]) + " |")
    lines += ["", "## Obserwacje", "", "| Forma | Przypadek | Lemat modelu | Nazwa kanoniczna PRNG | Proponowany lemat | Zmiana | Tokeny | Kategoria | Decyzja |", "|---|---|---|---|---|---|---:|---|---|"]
    for result in sorted(results, key=lambda item: (-item["token_count"], norm(item["orth_raw"]))):
        lines.append("| " + " | ".join([
            safe_md(result["orth_raw"]), safe_md(result["observed_case"]),
            safe_md(result["model_lemma"]), safe_md(result.get("canonical_name", "")),
            safe_md(result["proposed_lemma"]), "TAK" if result.get("changes_lemma") else "NIE",
            f"{result['token_count']:,}", safe_md(result["category"]),
            safe_md(result.get("decision", "")),
        ]) + " |")
    lines += ["", "## Interpretacja", "",
        "- `PRNG_SAFE_LEMMA_REPAIR` oznacza dokładne, jednoznaczne dopasowanie formy i przypadka oraz rzeczywistą zmianę lematu modelu.",
        "- `PRNG_MODEL_LEMMA_CONFIRMED` oznacza ten sam mocny dowód PRNG, ale lemat modelu był już poprawny i pozostaje bez zmiany.",
        "- `PRNG_ALTERNATIVE_NAME_CONFIRMED` identyfikuje nazwę oboczną lub endonim, lecz nie zastępuje lematu tokenu nazwą główną.",
        "- `PRNG_NAME_ONLY_CONFIRMED` potwierdza istnienie nazwy, ale bez zgodności kolumny przypadka; wynik pozostaje diagnostyczny.",
        "- Wynik niejednoznaczny nie jest kandydatem automatycznym.",
        "- Brak dopasowania nie jest dowodem niepoprawności nazwy.", "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8", newline="\n")


def run_audit(audit: Path, index: Path, report: Path, include_review: bool, min_count: int) -> dict[str, Any]:
    rows = load_candidates(audit, include_review, min_count)
    con = connect(index)
    try:
        results = [classify_candidate(row, query_name(con, normalize_display(row.get("orth")))) for row in rows]
    finally:
        con.close()
    rules = aggregate(results)
    report.parent.mkdir(parents=True, exist_ok=True)
    raw = report.with_suffix(".json")
    payload = {
        "schema_version": SCHEMA_VERSION, "version": VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "audit_json": str(audit.resolve()), "index": str(index.resolve()),
        "results": results, "rules": rules,
    }
    raw.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    write_report(report, audit.resolve(), index.resolve(), results, rules)
    return {
        "success": True, "experimental": True, "version": VERSION,
        "rows": len(results), "rules": len(rules),
        "report": str(report.resolve()), "json": str(raw.resolve()),
        "categories": dict(Counter(result["category"] for result in results)),
        "confidence": dict(Counter(result["confidence"] for result in results)),
        "safe_lemma_repairs": sum(bool(result.get("changes_lemma")) for result in results),
        "lemma_tokens_changed": sum(result["token_count"] for result in results if result.get("changes_lemma")),
        "model_lemma_confirmations": sum(result.get("category") == "PRNG_MODEL_LEMMA_CONFIRMED" for result in results),
        "confirmed_tokens_unchanged": sum(result["token_count"] for result in results if result.get("category") == "PRNG_MODEL_LEMMA_CONFIRMED"),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    build = sub.add_parser("build-index", help="Zbuduj lokalny indeks PRNG")
    build.add_argument("--source", required=True, type=Path, action="append", help="XLSX, CSV/TSV lub ZIP; opcje mozna powtarzac")
    build.add_argument("--index", required=True, type=Path)
    build.add_argument("--batch-size", type=int, default=5000)

    audit_parser = sub.add_parser("audit", help="Porownaj kandydatow z indeksem PRNG")
    audit_parser.add_argument("--audit-json", required=True, type=Path)
    audit_parser.add_argument("--index", required=True, type=Path)
    audit_parser.add_argument("--report", type=Path)
    audit_parser.add_argument("--include-review", action="store_true")
    audit_parser.add_argument("--min-count", type=int, default=2)

    args = parser.parse_args(argv)
    if args.command == "build-index":
        sources = [path.expanduser().resolve() for path in args.source]
        missing = [str(path) for path in sources if not path.is_file()]
        if missing:
            raise SystemExit("Brak plikow PRNG: " + ", ".join(missing))
        result = build_index(sources, args.index.expanduser().resolve(), args.batch_size)
    else:
        audit = args.audit_json.expanduser().resolve()
        index = args.index.expanduser().resolve()
        if not audit.is_file():
            raise SystemExit(f"Brak JSON audytu: {audit}")
        if not index.is_file():
            raise SystemExit(f"Brak indeksu PRNG: {index}")
        report = (args.report or audit.with_name(audit.stem + ".local_prng.md")).expanduser().resolve()
        result = run_audit(audit, index, report, args.include_review, args.min_count)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
